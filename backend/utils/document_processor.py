import os
import PyPDF2
import docx
from langchain.text_splitter import RecursiveCharacterTextSplitter
from db.database import SessionLocal, Document
from db.vector_store import store_document_chunks
from utils.embeddings import get_embeddings

def extract_text_from_pdf(file_path):
    """
    Extract text from a PDF file
    
    Args:
        file_path: Path to the PDF file
        
    Returns:
        Extracted text as a string
    """
    with open(file_path, 'rb') as file:
        pdf_reader = PyPDF2.PdfReader(file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text() + "\n"
    return text

def extract_text_from_docx(file_path):
    """
    Extract text from a DOCX file
    
    Args:
        file_path: Path to the DOCX file
        
    Returns:
        Extracted text as a string
    """
    doc = docx.Document(file_path)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

def extract_text_from_txt(file_path):
    """
    Extract text from a TXT file
    
    Args:
        file_path: Path to the TXT file
        
    Returns:
        Extracted text as a string
    """
    with open(file_path, 'r', encoding='utf-8') as file:
        text = file.read()
    return text

def extract_text(file_path):
    """
    Extract text from a file based on its extension
    
    Args:
        file_path: Path to the file
        
    Returns:
        Extracted text as a string
    """
    file_extension = os.path.splitext(file_path)[1].lower()
    
    if file_extension == '.pdf':
        return extract_text_from_pdf(file_path)
    elif file_extension == '.docx':
        return extract_text_from_docx(file_path)
    elif file_extension == '.txt':
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file extension: {file_extension}")

def chunk_text(text, filename):
    """
    Split text into chunks for processing
    
    Args:
        text: Text to split
        filename: Original filename for metadata
        
    Returns:
        List of chunks with content and metadata
    """
    # Create text splitter
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=200,
        length_function=len,
    )
    
    # Split text into chunks
    raw_chunks = text_splitter.split_text(text)
    
    # Add metadata to chunks
    chunks = []
    for i, chunk_text in enumerate(raw_chunks):
        chunks.append({
            "content": chunk_text,
            "metadata": {
                "source": filename,
                "chunk_index": i
            }
        })
    
    return chunks

def process_document(file_path):
    """
    Process a document: extract text, chunk it, generate embeddings, and store in database
    
    Args:
        file_path: Path to the document file
        
    Returns:
        Document ID
    """
    db = SessionLocal()
    try:
        # Extract filename
        filename = os.path.basename(file_path)
        
        # Create document record
        document = Document(
            filename=filename,
            title=os.path.splitext(filename)[0]  # Use filename without extension as title
        )
        db.add(document)
        db.commit()
        db.refresh(document)
        
        # Extract text from document
        text = extract_text(file_path)
        
        # Chunk text
        chunks = chunk_text(text, filename)
        
        # Generate embeddings for chunks
        chunk_texts = [chunk["content"] for chunk in chunks]
        embeddings = get_embeddings(chunk_texts)
        
        # Store chunks and embeddings
        store_document_chunks(document.id, chunks, embeddings)
        
        return document.id
    except Exception as e:
        db.rollback()
        raise e
    finally:
        db.close()
